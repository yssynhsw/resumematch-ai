"""Analysis orchestration: scoring, recommendations, cover letters."""

from __future__ import annotations

import logging
import threading
from datetime import timedelta
from io import BytesIO
from typing import Callable

from django.conf import settings
from django.db import close_old_connections, transaction
from django.utils import timezone

from cvs.models import CV
from services.ai_service import applicant_skill_names, get_ai_service, skill_mentioned_in
from services.ats_service import collect_pdf_layout, evaluate_ats
from services.scoring import (
    apply_interview_score,
    canonicalize_skills,
    skill_overlap_score,
)

from .models import Analysis, AnalysisJob, ApplicantInterview, CoverLetter, InterviewAttempt, JobOffer

logger = logging.getLogger(__name__)

ProgressFn = Callable[[int, str], None]


def _noop_progress(_pct: int, _step: str) -> None:
    return None


def run_analysis(
    *,
    user,
    cv: CV,
    job_offer_text: str,
    title: str = "",
    company: str = "",
    level: str = Analysis.Level.JUNIOR,
    language: str = "fr",
    progress: ProgressFn = _noop_progress,
) -> Analysis:
    ai = get_ai_service()

    progress(10, "cv")
    if not cv.parsed_data and cv.extracted_text:
        cv.parsed_data = ai.extract_cv_skills(cv.extracted_text).model_dump()
        cv.save(update_fields=["parsed_data"])

    progress(25, "job")
    job_parsed = ai.analyze_job_offer(job_offer_text)
    resolved_title = title.strip() or job_parsed.title or "Offre sans titre"
    resolved_company = company.strip() or job_parsed.company or ""

    job_offer = JobOffer.objects.create(
        user=user,
        title=resolved_title,
        company=resolved_company,
        raw_text=job_offer_text,
        parsed_data=job_parsed.model_dump(),
    )

    cv_skills = list(cv.parsed_data.get("technical_skills") or []) + list(
        cv.parsed_data.get("soft_skills") or []
    )
    required = list(job_parsed.required_skills or [])
    nice = list(job_parsed.nice_to_have_skills or [])

    skills_score, present, missing = skill_overlap_score(cv_skills, required)

    # Nice-to-have bonus (small)
    nice_present = 0
    if nice:
        _, nice_hit, _ = skill_overlap_score(cv_skills, nice)
        nice_present = len(nice_hit)

    progress(45, "semantic")
    semantic = ai.evaluate_semantic_fit(
        cv.extracted_text or "",
        job_offer_text,
        level=level,  # type: ignore[arg-type]
        present_skills=present,
        missing_skills=missing,
    )

    w_skills = settings.SCORE_WEIGHT_SKILLS
    w_semantic = settings.SCORE_WEIGHT_SEMANTIC
    level_mult = settings.SCORE_LEVEL_MULTIPLIERS.get(level, 1.0)

    # Renormalize weights after level multiplier on semantic part
    adj_semantic = w_semantic * level_mult
    total_w = w_skills + adj_semantic
    skills_contrib = (w_skills / total_w) * skills_score
    semantic_contrib = (adj_semantic / total_w) * semantic.semantic_score
    nice_bonus = min(5.0, nice_present * 1.5)
    final_score = int(round(min(100.0, skills_contrib + semantic_contrib + nice_bonus)))

    progress(70, "recommendations")
    recs = ai.generate_recommendations(
        cv.extracted_text or "",
        job_offer_text,
        level=level,  # type: ignore[arg-type]
        missing_skills=missing,
        detailed=False,
        language=language,  # type: ignore[arg-type]
    )
    progress(90, "saving")

    breakdown = {
        "skills_score": round(skills_score, 1),
        "semantic_score": semantic.semantic_score,
        "skills_weight": round(w_skills / total_w, 3),
        "semantic_weight": round(adj_semantic / total_w, 3),
        "skills_contribution": round(skills_contrib, 1),
        "semantic_contribution": round(semantic_contrib, 1),
        "nice_to_have_bonus": round(nice_bonus, 1),
        "level": level,
        "level_semantic_multiplier": level_mult,
        "rationale": semantic.rationale,
        "experience_fit": semantic.experience_fit,
        "present_count": len(present),
        "missing_count": len(missing),
        "required_count": len(canonicalize_skills(required)),
    }

    analysis = Analysis.objects.create(
        user=user,
        cv=cv,
        job_offer=job_offer,
        level=level,
        score=final_score,
        present_skills=present,
        missing_skills=missing,
        recommendations=recs.recommandations,
        score_breakdown=breakdown,
        status=Analysis.Status.TO_APPLY,
        language=language,
        ai_mode=ai.mode,
    )
    progress(100, "done")
    return analysis


def analyses_today(user) -> int:
    since = timezone.now() - timedelta(days=1)
    return Analysis.objects.filter(user=user, created_at__gte=since).count()


def effective_daily_quota() -> int:
    from adminpanel.models import SiteSettings

    return SiteSettings.load().effective_daily_quota


def quota_remaining(user, extra_pending: int = 0) -> int:
    limit = effective_daily_quota()
    if limit <= 0:
        return 10**6
    pending = AnalysisJob.objects.filter(
        user=user, state__in=[AnalysisJob.State.PENDING, AnalysisJob.State.RUNNING]
    ).count()
    return max(0, limit - analyses_today(user) - pending - extra_pending)


def start_analysis_job(*, user, cv: CV, **kwargs) -> AnalysisJob:
    job = AnalysisJob.objects.create(user=user)

    def update(pct: int, step: str) -> None:
        AnalysisJob.objects.filter(pk=job.pk).update(
            state=AnalysisJob.State.RUNNING, progress=pct, step=step, updated_at=timezone.now()
        )

    in_thread = not settings.ANALYSIS_JOBS_SYNC

    def worker() -> None:
        if in_thread:
            close_old_connections()
        try:
            analysis = run_analysis(user=user, cv=cv, progress=update, **kwargs)
            AnalysisJob.objects.filter(pk=job.pk).update(
                state=AnalysisJob.State.DONE,
                progress=100,
                step="done",
                analysis=analysis,
                updated_at=timezone.now(),
            )
        except Exception as exc:  # noqa: BLE001
            logger.exception("Analysis job %s failed", job.pk)
            AnalysisJob.objects.filter(pk=job.pk).update(
                state=AnalysisJob.State.FAILED,
                error=str(exc)[:500] or "Erreur inconnue",
                updated_at=timezone.now(),
            )
        finally:
            if in_thread:
                close_old_connections()

    if in_thread:
        threading.Thread(target=worker, daemon=True, name=f"analysis-job-{job.pk}").start()
    else:
        worker()
        job.refresh_from_db()
    return job


def _extras_key(kind: str, language: str) -> str:
    return f"{kind}:{language}"


def get_or_generate_extra(analysis: Analysis, kind: str, language: str, refresh: bool = False) -> dict:
    """Generate (or return cached) interview_prep / cv_rewrite / learning_plan / ats content."""
    key = _extras_key(kind, language)
    extras = dict(analysis.extras or {})
    if not refresh and key in extras:
        return extras[key]

    ai = get_ai_service()
    cv_text = analysis.cv.extracted_text or ""
    job_text = analysis.job_offer.raw_text
    missing = analysis.missing_skills or []

    if kind == "interview_prep":
        data = ai.generate_interview_prep(
            cv_text,
            job_text,
            analysis.level,  # type: ignore[arg-type]
            missing,
            language,  # type: ignore[arg-type]
            score=analysis.score,
            present_skills=list(analysis.present_skills or []),
        ).model_dump()
    elif kind == "cv_rewrite":
        data = ai.rewrite_cv(
            cv_text, job_text, analysis.level, missing, language  # type: ignore[arg-type]
        ).model_dump()
    elif kind == "learning_plan":
        data = ai.generate_learning_plan(
            missing, analysis.level, analysis.job_offer.title, language  # type: ignore[arg-type]
        ).model_dump()
    elif kind == "ats":
        layout = {}
        try:
            layout = collect_pdf_layout(analysis.cv.file.path)
        except Exception:  # noqa: BLE001
            logger.warning("CV file unavailable for ATS layout", exc_info=True)
        required = list((analysis.job_offer.parsed_data or {}).get("required_skills") or [])
        data = evaluate_ats(cv_text, layout, required, language)
    else:
        raise ValueError(f"Unknown extra kind: {kind}")

    if kind != "ats":
        data["ai_mode"] = ai.mode
    extras[key] = data
    analysis.extras = extras
    analysis.save(update_fields=["extras"])
    return data


def build_report_docx(analysis: Analysis, language: str = "fr") -> bytes:
    from docx import Document

    en = language == "en"
    t = (lambda fr, eng: eng if en else fr)
    breakdown = analysis.score_breakdown or {}
    doc = Document()
    doc.add_heading(t("Rapport d'analyse – ResumeMatch AI", "Analysis report – ResumeMatch AI"), 0)
    doc.add_paragraph(
        f"{analysis.job_offer.title} · {analysis.job_offer.company or '—'} · "
        f"{analysis.created_at:%d/%m/%Y}"
    )
    doc.add_heading(t(f"Score : {analysis.score}/100", f"Score: {analysis.score}/100"), 1)
    doc.add_paragraph(
        t("Compétences", "Skills") + f" : {breakdown.get('skills_score', '—')} · "
        + t("Sémantique", "Semantic") + f" : {breakdown.get('semantic_score', '—')}"
    )
    if breakdown.get("rationale"):
        doc.add_paragraph(str(breakdown["rationale"]))

    def bullets(title: str, items: list) -> None:
        doc.add_heading(title, 2)
        if not items:
            doc.add_paragraph("—")
        for item in items:
            doc.add_paragraph(str(item), style="List Bullet")

    bullets(t("Compétences présentes", "Matched skills"), analysis.present_skills or [])
    bullets(t("Compétences manquantes", "Missing skills"), analysis.missing_skills or [])
    bullets(t("Recommandations", "Recommendations"), analysis.recommendations or [])

    extras = analysis.extras or {}
    ats = extras.get(_extras_key("ats", language))
    if ats:
        doc.add_heading(t(f"Compatibilité ATS : {ats['score']}/100", f"ATS compatibility: {ats['score']}/100"), 2)
        for check in ats.get("checks", []):
            doc.add_paragraph(f"[{check['status'].upper()}] {check['label']} – {check['detail']}", style="List Bullet")

    rewrite = extras.get(_extras_key("cv_rewrite", language))
    if rewrite:
        doc.add_heading(t("Réécriture du CV", "CV rewrite"), 2)
        for item in rewrite.get("items", []):
            if item.get("before"):
                doc.add_paragraph(t("Avant : ", "Before: ") + item["before"])
            doc.add_paragraph(t("Après : ", "After: ") + item["after"])

    interview = extras.get(_extras_key("interview_prep", language))
    if interview:
        doc.add_heading(t("Préparation d'entretien", "Interview preparation"), 2)
        for q in interview.get("questions", []):
            doc.add_paragraph(q["question"], style="List Number")
            if q.get("answer_tips"):
                doc.add_paragraph(q["answer_tips"])

    plan = extras.get(_extras_key("learning_plan", language))
    if plan:
        doc.add_heading(t("Plan d'apprentissage", "Learning plan"), 2)
        for step in plan.get("steps", []):
            doc.add_paragraph(
                f"{step['skill']} – {step['priority']} – ~{step['estimated_weeks']} "
                + t("semaine(s)", "week(s)"),
                style="List Bullet",
            )
            if step.get("project_idea"):
                doc.add_paragraph(step["project_idea"])

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def improve_cv_suggestions(analysis: Analysis, language: str = "fr") -> dict:
    ai = get_ai_service()
    result = ai.generate_recommendations(
        analysis.cv.extracted_text or "",
        analysis.job_offer.raw_text,
        level=analysis.level,  # type: ignore[arg-type]
        missing_skills=analysis.missing_skills or [],
        detailed=True,
        language=language,  # type: ignore[arg-type]
    )
    return {
        "recommandations": result.recommandations,
        "detailed_improvements": result.detailed_improvements,
        "niveau": analysis.level,
        "analysis_id": analysis.id,
        "ai_mode": ai.mode,
    }


def create_cover_letter(analysis: Analysis, tone: str, language: str = "fr") -> CoverLetter:
    ai = get_ai_service()
    user = analysis.user
    name = (user.get_full_name() or user.username).strip()
    result = ai.generate_cover_letter(
        analysis.cv.extracted_text or "",
        analysis.job_offer.raw_text,
        level=analysis.level,  # type: ignore[arg-type]
        tone=tone,  # type: ignore[arg-type]
        candidate_name=name,
        job_title=analysis.job_offer.title,
        company=analysis.job_offer.company,
        language=language,  # type: ignore[arg-type]
    )
    return CoverLetter.objects.create(
        analysis=analysis,
        tone=tone,
        content=result.content,
    )


def personal_interview_questions(language: str = "fr") -> list[dict]:
    """Personal questions whose answers are reused in the application email."""
    en = language == "en"
    return [
        {
            "question": (
                "What work environment helps you do your best?"
                if en
                else "Dans quel environnement de travail êtes-vous le plus à l'aise ?"
            ),
            "category": "environnement",
            "why": (
                "Shows how you fit the team's way of working."
                if en
                else "Montre comment vous vous intégrez à la façon de travailler de l'équipe."
            ),
            "answer_tips": (
                "Describe pace, autonomy, feedback and the kind of team you prefer, using a real example."
                if en
                else "Décrivez le rythme, l'autonomie, le feedback et l'équipe que vous préférez, avec un exemple réel."
            ),
        },
        {
            "question": (
                "Which soft skill do you use most with a team, and how?"
                if en
                else "Quel soft skill utilisez-vous le plus en équipe, et comment ?"
            ),
            "category": "comportementale",
            "why": (
                "Looks for a concrete personal quality, not a slogan."
                if en
                else "Cherche une qualité personnelle concrète, pas un slogan."
            ),
            "answer_tips": (
                "Pick one skill such as listening, organisation or collaboration, and tell a short true story."
                if en
                else "Choisissez une qualité comme l'écoute, l'organisation ou la collaboration, et racontez un fait vrai."
            ),
        },
    ]


def _with_personal_questions(questions: list, language: str) -> list:
    current = list(questions or [])
    known = {str(item.get("question") or "").strip() for item in current}
    for question in personal_interview_questions(language):
        if question["question"] not in known:
            current.append(question)
            known.add(question["question"])
    return current


def _cv_skill_lists(cv: CV) -> tuple[list[str], list[str]]:
    parsed = cv.parsed_data or {}

    def clean(values) -> list[str]:
        seen = set()
        result = []
        for value in values or []:
            text = str(value or "").strip()
            key = text.lower()
            if text and key not in seen:
                seen.add(key)
                result.append(text)
        return result

    return clean(parsed.get("technical_skills")), clean(parsed.get("soft_skills"))


def _work_on_summary(question_scores: list[dict], language: str = "fr") -> str:
    """The weakest answers first: what the candidate should practice next."""
    ranked = sorted(question_scores, key=lambda item: item.get("score") or 0)
    lines = []
    seen = set()
    for item in ranked:
        if (item.get("score") or 0) >= 75 and lines:
            continue
        text = str(item.get("focus") or "").strip()
        key = text.lower()
        if not text or key in seen:
            continue
        seen.add(key)
        lines.append(text)
        if len(lines) >= 4:
            break
    if not lines:
        return f"{round(sum(item.get('score') or 0 for item in question_scores) / len(question_scores))}/100" if question_scores else ""
    heading = "Work on:" if language == "en" else "À travailler :"
    return heading + "\n" + "\n".join(f"- {line}" for line in lines)


def _score_answers(questions: list, answers: list[str], language: str = "fr") -> tuple[list[str], list[dict], int]:
    cleaned = [str(answer or "").strip() for answer in answers]
    if len(cleaned) < len(questions):
        cleaned.extend([""] * (len(questions) - len(cleaned)))
    cleaned = cleaned[: len(questions)]
    graded = get_ai_service().score_interview_answers(questions, cleaned, language)
    question_scores = [
        {
            "index": item.index,
            "score": item.score,
            "question": (questions[item.index] or {}).get("question", "") if item.index < len(questions) else "",
            "comment": item.comment,
            "focus": item.focus,
        }
        for item in graded.scores
    ]
    overall = round(sum(item["score"] for item in question_scores) / len(question_scores)) if question_scores else 0
    return cleaned, question_scores, overall


_SOFT_CATEGORIES = {"comportementale", "environnement"}


def _separate_legacy_mixed_interview(cv: CV) -> None:
    """Move soft-skill questions out of an older mixed interview into the soft type."""
    with transaction.atomic():
        technical = (
            ApplicantInterview.objects.select_for_update()
            .filter(cv=cv, kind=ApplicantInterview.Kind.TECHNICAL)
            .first()
        )
        if technical is None:
            return
        if ApplicantInterview.objects.filter(cv=cv, kind=ApplicantInterview.Kind.SOFT).exists():
            return
        questions = list(technical.questions or [])
        if not any(item.get("category") in _SOFT_CATEGORIES for item in questions):
            return
        answers = list(technical.answers or [])
        stored_scores = {item.get("index"): item.get("score") for item in (technical.question_scores or [])}
        grouped = {ApplicantInterview.Kind.TECHNICAL: [], ApplicantInterview.Kind.SOFT: []}
        for index, question in enumerate(questions):
            target = (
                ApplicantInterview.Kind.SOFT
                if question.get("category") in _SOFT_CATEGORIES
                else ApplicantInterview.Kind.TECHNICAL
            )
            grouped[target].append((question, answers[index] if index < len(answers) else "", stored_scores.get(index)))
        for kind, rows in grouped.items():
            if not rows:
                continue
            interview = technical if kind == ApplicantInterview.Kind.TECHNICAL else ApplicantInterview(
                cv=cv, user=cv.user, kind=kind
            )
            interview.kind = kind
            interview.questions = [row[0] for row in rows]
            interview.answers = [row[1] for row in rows]
            interview.question_scores = [
                {"index": index, "score": row[2] or 0, "question": row[0].get("question", "")}
                for index, row in enumerate(rows)
                if row[2] is not None
            ]
            if technical.state == ApplicantInterview.State.SCORED and interview.question_scores:
                interview.score = round(
                    sum(item["score"] for item in interview.question_scores) / len(interview.question_scores)
                )
                interview.summary = f"{interview.score}/100"
                interview.state = ApplicantInterview.State.SCORED
            if interview.kind == ApplicantInterview.Kind.TECHNICAL and interview.state == ApplicantInterview.State.SCORED:
                interview.courses = _courses_for_scored_interview(interview, "fr")
            else:
                interview.courses = []
            interview.save()


def _skills_for_kind(cv: CV, kind: str) -> list[str]:
    technical, soft = _cv_skill_lists(cv)
    source = soft if kind == ApplicantInterview.Kind.SOFT else technical
    return applicant_skill_names(source)


def _questions_match_skills(questions: list, skills: list[str], kind: str) -> bool:
    """True when stored questions are exactly one per CV skill (or one fallback)."""
    category = "comportementale" if kind == ApplicantInterview.Kind.SOFT else "technique"
    items = [item for item in (questions or []) if isinstance(item, dict)]
    selected = applicant_skill_names(skills)
    if not selected:
        return len(items) == 1 and items[0].get("category") == category
    if len(items) != len(selected):
        return False
    used: set[int] = set()
    for skill in sorted(selected, key=len, reverse=True):
        found = None
        for index, item in enumerate(items):
            if index in used or item.get("category") != category:
                continue
            if skill_mentioned_in(str(item.get("question") or ""), skill):
                found = index
                break
        if found is None:
            return False
        used.add(found)
    return len(used) == len(items)


def _update_user_interview_score(user):
    """Rounded average of this user's scored applicant interviews, or null."""
    scores = list(
        ApplicantInterview.objects.filter(
            user=user,
            state=ApplicantInterview.State.SCORED,
            score__isnull=False,
        ).values_list("score", flat=True)
    )
    user.score = round(sum(scores) / len(scores)) if scores else None
    user.save(update_fields=["score"])
    return user.score


def _question_points(item: dict) -> int:
    try:
        return max(0, min(100, int(round(float(item.get("score") or 0)))))
    except (TypeError, ValueError):
        return 0


def _scored_question_text(interview: ApplicantInterview, item: dict, position: int) -> str:
    text = str(item.get("question") or "").strip()
    if text:
        return text
    questions = list(interview.questions or [])
    try:
        index = int(item.get("index", position))
    except (TypeError, ValueError):
        index = position
    if 0 <= index < len(questions) and isinstance(questions[index], dict):
        return str(questions[index].get("question") or "")
    return ""


def _skills_for_courses(interview: ApplicantInterview) -> list[str]:
    """Up to four CV technical skills named in the scored questions.

    A skill whose question score is under 80 is preferred. When every matched
    score is 80 or higher, the lowest scores are kept.
    """
    technical, _soft = _cv_skill_lists(interview.cv)
    skills = applicant_skill_names(technical)
    ranked: list[tuple[int, int, str]] = []
    for order, skill in enumerate(skills):
        scores: list[int] = []
        for position, item in enumerate(interview.question_scores or []):
            if not isinstance(item, dict):
                continue
            question = _scored_question_text(interview, item, position)
            if question and skill_mentioned_in(question, skill):
                scores.append(_question_points(item))
        if scores:
            ranked.append((min(scores), order, skill))
    weak = [row for row in ranked if row[0] < 80]
    pool = weak or ranked
    pool.sort(key=lambda row: (row[0], row[1]))
    return [skill for _score, _order, skill in pool[:4]]


def _courses_for_scored_interview(interview: ApplicantInterview, language: str = "fr") -> list[dict]:
    skills = _skills_for_courses(interview)
    if not skills:
        return []
    language = "en" if language == "en" else "fr"
    return list(get_ai_service().generate_interview_courses(skills, language) or [])


def _backfill_courses(interview: ApplicantInterview, language: str) -> bool:
    """Fill an already scored technical interview when its course list is empty."""
    if interview.kind != ApplicantInterview.Kind.TECHNICAL:
        return False
    if interview.state != ApplicantInterview.State.SCORED:
        return False
    if interview.courses:
        return False
    interview.courses = _courses_for_scored_interview(interview, language)
    return bool(interview.courses)


def start_applicant_interview(cv: CV, language: str = "fr", kind: str = "technical") -> ApplicantInterview:
    """Build an evaluation from the uploaded CV skills only. No job application is involved."""
    _separate_legacy_mixed_interview(cv)
    if kind not in (ApplicantInterview.Kind.TECHNICAL, ApplicantInterview.Kind.SOFT):
        kind = ApplicantInterview.Kind.TECHNICAL
    skills = _skills_for_kind(cv, kind)
    existing = ApplicantInterview.objects.filter(cv=cv, kind=kind).first()
    if (
        existing is not None
        and existing.state == ApplicantInterview.State.SCORED
        and _questions_match_skills(existing.questions, skills, kind)
    ):
        if _backfill_courses(existing, language):
            existing.save(update_fields=["courses", "updated_at"])
        return existing
    technical, soft = _cv_skill_lists(cv)
    prepared = get_ai_service().generate_applicant_interview(
        cv.extracted_text or "",
        technical,
        soft,
        language,  # type: ignore[arg-type]
        kind,
    )
    questions = prepared.model_dump().get("questions") or []
    interview, _created = ApplicantInterview.objects.get_or_create(
        cv=cv,
        kind=kind,
        defaults={"user": cv.user},
    )
    if interview.user_id != cv.user_id:
        interview.user = cv.user
    if (
        interview.state == ApplicantInterview.State.SCORED
        and _questions_match_skills(interview.questions, skills, kind)
    ):
        _backfill_courses(interview, language)
        interview.save()
        return interview
    was_scored = interview.state == ApplicantInterview.State.SCORED
    interview.questions = questions
    interview.answers = []
    interview.question_scores = []
    interview.score = None
    interview.summary = ""
    interview.courses = []
    interview.state = ApplicantInterview.State.IN_PROGRESS
    interview.save()
    if was_scored:
        _update_user_interview_score(interview.user)
    return interview


def submit_applicant_interview(
    cv: CV, answers: list[str], kind: str = "technical", language: str = "fr"
) -> ApplicantInterview:
    """Score the answers and store the result on the interview and the user's main score."""
    interview = ApplicantInterview.objects.get(cv=cv, kind=kind)
    if interview.user_id != cv.user_id:
        interview.user = cv.user
    questions = list(interview.questions or [])
    cleaned, question_scores, overall = _score_answers(questions, answers, language)
    interview.answers = cleaned
    interview.question_scores = question_scores
    interview.score = overall
    interview.summary = _work_on_summary(question_scores, language)
    interview.state = ApplicantInterview.State.SCORED
    if interview.kind == ApplicantInterview.Kind.TECHNICAL:
        interview.courses = _courses_for_scored_interview(interview, language)
    else:
        interview.courses = []
    interview.save()
    _update_user_interview_score(interview.user)
    return interview


def start_interview_attempt(analysis: Analysis, language: str = "fr") -> InterviewAttempt:
    """Build questions from the CV and the offer. Tips stay stored for scoring only."""
    prepared = get_or_generate_extra(analysis, "interview_prep", language)
    questions = _with_personal_questions(prepared.get("questions") or [], language)
    attempt, _created = InterviewAttempt.objects.get_or_create(analysis=analysis)
    if attempt.state == InterviewAttempt.State.SCORED and attempt.questions:
        return attempt
    attempt.questions = questions
    attempt.answers = []
    attempt.question_scores = []
    attempt.score = None
    attempt.summary = ""
    attempt.state = InterviewAttempt.State.IN_PROGRESS
    attempt.save()
    return attempt


def extend_interview_with_personal_questions(analysis: Analysis, language: str = "fr") -> InterviewAttempt:
    """Add the personal questions to an existing interview and reopen it for those answers."""
    attempt = InterviewAttempt.objects.filter(analysis=analysis).first()
    if attempt is None:
        return start_interview_attempt(analysis, language)
    previous = list(attempt.questions or [])
    questions = _with_personal_questions(previous, language)
    answers = list(attempt.answers or [])
    if len(answers) < len(questions):
        answers.extend([""] * (len(questions) - len(answers)))
    added = len(questions) > len(previous)
    attempt.questions = questions
    attempt.answers = answers[: len(questions)]
    if added:
        attempt.state = InterviewAttempt.State.IN_PROGRESS
        attempt.question_scores = []
        attempt.score = None
        attempt.summary = ""
    attempt.save()
    return attempt


def submit_interview_attempt(analysis: Analysis, answers: list[str]) -> tuple[InterviewAttempt, Analysis]:
    attempt = analysis.interview_attempt
    questions = list(attempt.questions or [])
    cleaned = [str(answer or "").strip() for answer in answers]
    if len(cleaned) < len(questions):
        cleaned.extend([""] * (len(questions) - len(cleaned)))
    cleaned = cleaned[: len(questions)]

    cleaned, question_scores, overall = _score_answers(questions, cleaned, analysis.language or "fr")
    final, breakdown = apply_interview_score(analysis.score_breakdown, overall)

    attempt.answers = cleaned
    attempt.question_scores = question_scores
    attempt.score = overall
    attempt.summary = _work_on_summary(question_scores, analysis.language or "fr")
    attempt.state = InterviewAttempt.State.SCORED
    attempt.save()

    analysis.score = final
    analysis.score_breakdown = breakdown
    analysis.save(update_fields=["score", "score_breakdown"])
    return attempt, analysis
